import pdfplumber
import docx
import openpyxl

def read_file(file_path):
    text = ""

    if file_path.endswith(".pdf"):
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ""

    elif file_path.endswith(".docx"):
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + " "

    elif file_path.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()

    elif file_path.endswith(".xlsx"):
        wb = openpyxl.load_workbook(file_path)
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell:
                        text += str(cell) + " "

    return text.strip()
